import os
import re
import logging
from io import BytesIO

import requests
from flask import Flask, request, abort, jsonify
from openpyxl import load_workbook

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("harbi-bot")

BOT_TOKEN = os.environ.get("BOT_TOKEN", "").strip()
WEBHOOK_SECRET = os.environ.get("WEBHOOK_SECRET", "").strip()
APP_URL = os.environ.get("RENDER_EXTERNAL_URL", "").strip().rstrip("/")
API = f"https://api.telegram.org/bot{BOT_TOKEN}"
MAX_EXCEL_BYTES = 15 * 1024 * 1024

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_EXCEL_BYTES

# Render's local filesystem and process memory are not durable. The latest
# workbook is cached in memory until the service restarts or redeploys.
PRODUCTS = []
PRODUCT_SOURCE_NAME = ""

ABOUT = (
    "🌟 کافی‌نت و لوازم‌التحریر حربی\n\n"
    "💻 خدمات کافی‌نت: تایپ، پرینت، کپی، اسکن و خدمات اینترنتی.\n"
    "✏️ لوازم‌التحریر: انواع دفتر، خودکار، مداد و نوشت‌افزار دانش‌آموزی و اداری.\n\n"
    "❤️ رضایت شما، افتخار ماست."
)
SERVICES = (
    "💻 خدمات کافی‌نت حربی\n\n"
    "• تایپ و آماده‌سازی فایل\n• پرینت و چاپ مدارک\n• کپی و اسکن\n"
    "• راهنمایی برای خدمات اینترنتی و ثبت‌نام‌های آنلاین\n\n"
    "برای هماهنگی درباره خدمت موردنظر، پیام خود را همین‌جا بنویسید."
)
STATIONERY = (
    "✏️ لوازم‌التحریر حربی\n\n"
    "نام کالای موردنظر را بفرستید تا در فایل اکسل بارگذاری‌شده جست‌وجو شود.\n"
    "برای بارگذاری فهرست کالاها، فایل اکسل (.xlsx) را به‌صورت فایل در همین ربات ارسال کنید.\n"
    "قیمت یا موجودی خالی حدس زده نمی‌شود."
)
CONTACT = "📞 برای ارتباط با کافی‌نت و لوازم‌التحریر حربی، درخواست خود را در همین گفت‌وگو ارسال کنید."

MAIN_KEYBOARD = {
    "keyboard": [
        [{"text": "ℹ️ درباره ما"}, {"text": "💻 خدمات کافی‌نت"}],
        [{"text": "✏️ لوازم‌التحریر"}, {"text": "📞 تماس با ما"}]
    ],
    "resize_keyboard": True,
    "is_persistent": True
}


def telegram(method, payload):
    if not BOT_TOKEN:
        log.error("BOT_TOKEN environment variable is missing")
        return None
    try:
        response = requests.post(f"{API}/{method}", json=payload, timeout=20)
        response.raise_for_status()
        data = response.json()
        if not data.get("ok"):
            log.error("Telegram API error for %s: %s", method, data)
        return data
    except requests.RequestException:
        log.exception("Telegram API request failed for %s", method)
        return None


def send_message(chat_id, text, keyboard=True):
    payload = {"chat_id": chat_id, "text": text}
    if keyboard:
        payload["reply_markup"] = MAIN_KEYBOARD
    telegram("sendMessage", payload)


def normalize(value):
    text = str(value or "").strip().lower()
    text = text.translate(str.maketrans({
        "ي": "ی", "ى": "ی", "ك": "ک",
        "۰": "0", "۱": "1", "۲": "2", "۳": "3", "۴": "4",
        "۵": "5", "۶": "6", "۷": "7", "۸": "8", "۹": "9",
        "٠": "0", "١": "1", "٢": "2", "٣": "3", "٤": "4",
        "٥": "5", "٦": "6", "٧": "7", "٨": "8", "٩": "9",
    }))
    text = re.sub(r"[\u064b-\u065f\u0670\u0640]", "", text)
    text = re.sub(r"[^\w\s.\-/]", " ", text, flags=re.UNICODE)
    return " ".join(text.split())


def clean_header(value):
    return normalize(value).replace(" ", "")


def parse_workbook(file_bytes):
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    parsed = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [clean_header(cell) for cell in header_row]
            name_idx = next((i for i, h in enumerate(headers) if h in ("نام", "نامکالا", "کالا", "عنوان")), None)
            if name_idx is None:
                continue

            def col_index(candidates):
                return next((i for i, h in enumerate(headers) if h in candidates), None)

            code_idx = col_index({"کد", "کدکالا", "شناسه"})
            category_idx = col_index({"دسته", "دستهبندی", "گروه"})
            buy_idx = col_index({"قیمتخرید", "بهایخرید"})
            sell_idx = col_index({"قیمتفروش", "قیمتخردهفروشی", "قیمت"})
            stock_idx = col_index({"موجودی", "تعدادموجودی", "موجودیفعلی"})
            unit_idx = col_index({"واحد"})
            description_idx = col_index({"توضیحات", "شرح"})

            for row in rows:
                if name_idx >= len(row) or row[name_idx] is None:
                    continue
                name = str(row[name_idx]).strip()
                if not name:
                    continue

                def value_at(idx):
                    if idx is None or idx >= len(row):
                        return None
                    value = row[idx]
                    if value is None or (isinstance(value, str) and not value.strip()):
                        return None
                    return value

                parsed.append({
                    "name": name,
                    "normalized_name": normalize(name),
                    "code": value_at(code_idx),
                    "category": value_at(category_idx),
                    "buy_price": value_at(buy_idx),
                    "sell_price": value_at(sell_idx),
                    "stock": value_at(stock_idx),
                    "unit": value_at(unit_idx),
                    "description": value_at(description_idx),
                })
            if parsed:
                break
    finally:
        workbook.close()
    if not parsed:
        raise ValueError("ستون «نام» یا «نام کالا» در ردیف اول فایل پیدا نشد.")
    return parsed


def format_value(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return "ثبت نشده"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"{value:,}" if isinstance(value, (int, float)) else str(value)



def online_price_search(product_name):
    """Search public Iranian retail listings; never infer a price from missing data."""
    from html import unescape
    from html.parser import HTMLParser

    class ResultParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.results = []
            self.current = None
            self.capture = None

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            classes = attrs.get("class", "").split()
            if tag == "a" and "result__a" in classes:
                self.current = {"title": "", "url": attrs.get("href", ""), "snippet": ""}
                self.capture = "title"
            elif self.current is not None and tag in ("a", "div", "span") and "result__snippet" in classes:
                self.capture = "snippet"

        def handle_data(self, data):
            if self.current is not None and self.capture in ("title", "snippet"):
                key = self.capture
                self.current[key] = (self.current.get(key, "") + " " + data).strip()

        def handle_endtag(self, tag):
            if tag == "a" and self.capture == "title" and self.current is not None:
                self.results.append(self.current)
                self.current = None
                self.capture = None
            elif tag in ("div", "span") and self.capture == "snippet":
                self.capture = None

    query = f'"{product_name}" قیمت تومان (site:torob.com OR site:emalls.ir OR site:digikala.com)'
    try:
        response = requests.get(
            "https://html.duckduckgo.com/html/",
            params={"q": query},
            headers={"User-Agent": "Mozilla/5.0 (compatible; HarbiShopPriceLookup/1.0)"},
            timeout=8,
        )
        response.raise_for_status()
        parser = ResultParser()
        parser.feed(response.text)
        results = []
        seen = set()
        for item in parser.results:
            url = unescape(item.get("url", "")).strip()
            title = unescape(item.get("title", "")).strip()
            snippet = unescape(item.get("snippet", "")).strip()
            if not url or not title or url in seen:
                continue
            seen.add(url)
            if not any(domain in url.lower() for domain in ("torob.com", "emalls.ir", "digikala.com")):
                continue
            match = re.search(r"([0-9۰-۹٠-٩][0-9۰-۹٠-٩,٬.\s]{1,18})\s*(تومان|تومن|ریال)", f"{title} {snippet}", re.IGNORECASE)
            price = None
            if match:
                digits = re.sub(r"[^0-9۰-۹٠-٩]", "", match.group(1))
                digits = digits.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
                if digits:
                    price = (digits, match.group(2))
            results.append({"title": title[:180], "url": url, "snippet": snippet[:220], "price": price})
            if len(results) >= 3:
                break
        return results
    except (requests.RequestException, ValueError) as exc:
        log.info("Online price search unavailable: %s", exc)
        return None


def product_reply(query):
    if not PRODUCTS:
        return (
            "هنوز فایل اکسل کالاها بارگذاری نشده است.\n"
            "لطفاً فایل .xlsx را به‌صورت فایل (Document) برای ربات بفرستید؛ "
            "سپس نام کالا را ارسال کنید."
        )

    q = normalize(query)
    if not q:
        return "نام کالا را ارسال کنید."
    exact = [p for p in PRODUCTS if p["normalized_name"] == q or (p["code"] is not None and normalize(p["code"]) == q)]
    matches = exact or [p for p in PRODUCTS if q in p["normalized_name"] or p["normalized_name"] in q]
    if not matches:
        return f"کالای «{query}» در فایل اکسل پیدا نشد. نام کالا را دقیق‌تر بفرستید یا فایل به‌روز را بارگذاری کنید."
    if len(matches) > 1:
        lines = ["چند کالا با این عبارت پیدا شد. نام دقیق‌تر را بفرستید:"]
        for p in matches[:8]:
            lines.append(f"• {p['name']}" + (f" (کد: {p['code']})" if p["code"] is not None else ""))
        if len(matches) > 8:
            lines.append("نتایج بیشتری هم وجود دارد.")
        return "\n".join(lines)

    p = matches[0]
    lines = [f"📦 {p['name']}"]
    if p["code"] is not None:
        lines.append(f"کد کالا: {p['code']}")
    if p["category"] is not None:
        lines.append(f"دسته: {p['category']}")
    lines.append(f"قیمت فروش ثبت‌شده: {format_value(p['sell_price'])}" + ("" if p["sell_price"] is None else " تومان"))
    lines.append(f"موجودی ثبت‌شده: {format_value(p['stock'])}" + (f" {p['unit']}" if p['stock'] is not None and p['unit'] is not None else ""))
    if p["buy_price"] is not None:
        lines.append(f"قیمت خرید ثبت‌شده: {format_value(p['buy_price'])} تومان")
    lines.append("\\n🌐 بررسی قیمت آنلاین از فروشگاه‌های ایرانی:")
    online = online_price_search(p["name"])
    if online is None:
        lines.append("جست‌وجوی وب در دسترس نیست؛ قیمت آنلاین نامشخص است.")
    elif not online:
        lines.append("نتیجه قابل استفاده‌ای پیدا نشد؛ قیمت آنلاین نامشخص است.")
    else:
        for item in online:
            lines.append(f"• {item['title']}")
            if item["price"]:
                lines.append(f"  قیمت درج‌شده در نتیجه جست‌وجو: {format_value(int(item['price'][0]))} {item['price'][1]} (نیازمند تأیید در صفحه فروشنده)")
            else:
                lines.append("  قیمت قابل تأیید از متن نتیجه پیدا نشد.")
            if item["snippet"]:
                lines.append(f"  {item['snippet']}")
            lines.append(f"  لینک: {item['url']}")
    lines.append("\\n⚠️ قیمت آنلاین ممکن است قدیمی یا مربوط به مدل متفاوت باشد. پیش از خرید صفحه فروشنده را بررسی کنید. قیمت و موجودی اکسل تغییر نمی‌کند و هیچ مقدار نامشخصی حدس زده نمی‌شود.")
    return "\\n".join(lines)


def download_telegram_file(file_id):
    info = telegram("getFile", {"file_id": file_id})
    if not info or not info.get("ok"):
        raise ValueError("دریافت مشخصات فایل از تلگرام ناموفق بود.")
    file_path = info.get("result", {}).get("file_path")
    if not file_path:
        raise ValueError("مسیر فایل در پاسخ تلگرام موجود نیست.")
    url = f"https://api.telegram.org/file/bot{BOT_TOKEN}/{file_path}"
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    if len(response.content) > MAX_EXCEL_BYTES:
        raise ValueError("حجم فایل بیشتر از ۱۵ مگابایت است.")
    return response.content


def handle_document(message):
    global PRODUCTS, PRODUCT_SOURCE_NAME
    chat_id = message["chat"]["id"]
    document = message.get("document") or {}
    filename = document.get("file_name") or "file"
    mime = (document.get("mime_type") or "").lower()
    if not filename.lower().endswith(".xlsx"):
        send_message(chat_id, "لطفاً فایل اکسل با پسوند .xlsx ارسال کنید. فایل CSV یا عکس قابل پردازش نیست.")
        return
    if document.get("file_size", 0) > MAX_EXCEL_BYTES:
        send_message(chat_id, "حجم فایل بیشتر از ۱۵ مگابایت است. نسخه کوچک‌تری ارسال کنید.")
        return
    try:
        content = download_telegram_file(document["file_id"])
        products = parse_workbook(content)
        PRODUCTS = products
        PRODUCT_SOURCE_NAME = filename
        send_message(
            chat_id,
            f"✅ فایل «{filename}» خوانده شد.\n"
            f"تعداد کالاهای شناسایی‌شده: {len(PRODUCTS)}\n\n"
            "حالا نام کالا را بفرستید تا قیمت و موجودی ثبت‌شده در اکسل را نمایش دهم. "
            "مقادیر خالی را حدس نمی‌زنم."
        )
    except Exception as exc:
        log.exception("Excel import failed")
        send_message(chat_id, f"خواندن فایل ناموفق بود: {str(exc)[:250]}\nلطفاً فایل .xlsx سالم و دارای ستون «نام» را ارسال کنید.")


def handle_update(update):
    message = update.get("message") or update.get("edited_message")
    if not message or not message.get("chat"):
        return
    chat_id = message["chat"]["id"]

    if message.get("document"):
        handle_document(message)
        return

    text = (message.get("text") or "").strip()
    if not text:
        send_message(chat_id, "لطفاً نام کالا را به‌صورت متن یا فایل اکسل .xlsx را ارسال کنید.")
        return

    normalized = normalize(text).split("@")[0]
    if normalized in ("/start", "/start start", "شروع"):
        send_message(chat_id, "سلام و خوش آمدید 🌷\nبه ربات کافی‌نت و لوازم‌التحریر حربی خوش آمدید.\nاز منوی زیر یکی از گزینه‌ها را انتخاب کنید.")
    elif normalized in ("/help", "راهنما"):
        send_message(chat_id, "از منوی زیر استفاده کنید. برای جست‌وجوی کالا، نام آن را بفرستید. برای بارگذاری یا به‌روزرسانی فهرست، فایل اکسل .xlsx را به ربات ارسال کنید.")
    elif normalized in ("/about", "ℹ️ درباره ما", "درباره ما", "درباره"):
        send_message(chat_id, ABOUT)
    elif normalized in ("/services", "💻 خدمات کافی‌نت", "خدمات کافی‌نت", "خدمات"):
        send_message(chat_id, SERVICES)
    elif normalized in ("/stationery", "✏️ لوازم‌التحریر", "لوازم‌التحریر", "لوازم التحریر"):
        send_message(chat_id, STATIONERY)
    elif normalized in ("/contact", "📞 تماس با ما", "تماس با ما", "ارتباط"):
        send_message(chat_id, CONTACT)
    else:
        send_message(chat_id, product_reply(text))


@app.get("/")
def home():
    return jsonify({"status": "ok", "service": "Harbi Telegram Bot", "products_loaded": len(PRODUCTS), "excel_source": PRODUCT_SOURCE_NAME or None})


@app.get("/health")
def health():
    return jsonify({"status": "healthy", "bot_token_configured": bool(BOT_TOKEN), "products_loaded": len(PRODUCTS)})


@app.post("/telegram-webhook")
def webhook():
    if WEBHOOK_SECRET:
        supplied = request.headers.get("X-Telegram-Bot-Api-Secret-Token", "")
        if supplied != WEBHOOK_SECRET:
            abort(403)
    update = request.get_json(silent=True)
    if not isinstance(update, dict):
        return jsonify({"ok": False, "error": "invalid update"}), 400
    handle_update(update)
    return jsonify({"ok": True})


if BOT_TOKEN and APP_URL and WEBHOOK_SECRET:
    result = telegram("setWebhook", {
        "url": f"{APP_URL}/telegram-webhook",
        "secret_token": WEBHOOK_SECRET,
        "allowed_updates": ["message"]
    })
    log.info("Webhook setup response: %s", result)
else:
    log.warning("Webhook not configured yet. Set BOT_TOKEN, WEBHOOK_SECRET, and deploy on Render.")
